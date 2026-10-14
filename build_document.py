from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
doc = Document()
s = doc.sections[0]
s.page_height = Cm(29.7); s.page_width = Cm(21)
s.top_margin = s.bottom_margin = Cm(1.8)
s.left_margin = s.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'; st.font.color.rgb = RGBColor(0,0,0)
for style in doc.styles:
    pr = style.element.find(qn('w:pPr'))
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.core_properties.title = 'AvecLesFemmes Résumé du besoin et plan de mise en œuvre'
doc.core_properties.author = 'Projet AvecLesFemmes'
f = s.footer.paragraphs[0]; f.alignment = 2
f.add_run('AvecLesFemmes  •  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,2)
def page(t): doc.add_page_break(); doc.add_heading(t,1)
def bullet(t): doc.add_paragraph(t, style='List Bullet')
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i,v in enumerate(headers): t.rows[0].cells[i].text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=v
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trpr.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DBE5F1' if ri==0 else 'FFFFFF'); pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs: run.font.size=Pt(10); run.bold=ri==0
    p('')

doc.add_paragraph('AvecLesFemmes', 'Title')
doc.add_paragraph('Résumé du besoin et plan de mise en œuvre', 'Subtitle')
p('Premier incrément • Document de cadrage proposé • Version 0.2 • 14 octobre 2026')
h('Objectif du projet')
p('AvecLesFemmes permet à des groupements de femmes d’organiser leur épargne solidaire, de gérer leurs campagnes et de suivre les versements des participantes. Le premier incrément couvre la gestion des groupements, des membres et des campagnes Avec Épargne, ainsi que les cotisations et les accès par rôle.')
p('Ce document sert de base au cadrage et à la réalisation. Les besoins confirmés sont distingués des recommandations techniques et des règles à valider. Les ADR restent au statut Proposé jusqu’à leur validation.')
h('Deux interfaces et un backend commun')
bullet('Application web pour les super administrateurs : gestion des groupements et des types de campagne, consultation globale et activation ou désactivation des campagnes.')
bullet('PWA pour les membres et les administrateurs de groupement : consultation personnelle et fonctions de gestion selon les permissions. Un administrateur est aussi une membre et utilise un seul compte.')
bullet('Backend commun : règles métier, autorisations, enregistrement des cotisations et accès aux données.')
h('Dimensionnement à anticiper')
p('La cible dépasse 1 000 groupements, chacun comptant plus de 200 membres. Une base de calcul de 1 000 × 200 représente 200 000 adhésions. Le nombre de personnes distinctes dépendra de la possibilité d’appartenir à plusieurs groupements.')
p('À raison d’un versement hebdomadaire par participante dans une seule campagne pendant 52 semaines, cette base représente environ 10,4 millions de cotisations par an. Ce scénario sert au dimensionnement ; il ne constitue pas une règle métier. Les pics de connexion et de collecte doivent encore être estimés.')
h('Résultat attendu du premier incrément')
p('Un groupement peut être créé, doté de membres et d’une campagne Avec, puis enregistrer des cotisations et consulter des totaux fiables. Chaque utilisatrice accède uniquement aux fonctions et données autorisées.')

page('1 Périmètre fonctionnel et droits')
h('Besoins confirmés')
bullet('Créer des groupements avec nom, adresse, responsables, téléphone et autres informations à préciser.')
bullet('Ajouter des membres aux groupements et consulter leurs listes.')
bullet('Créer des types de campagne et créer des campagnes par type au sein des groupements.')
bullet('Définir la date de début, la date de fin et la périodicité d’une campagne ; paramétrer le jour de cotisation, généralement un jour de la semaine.')
bullet('Pour Avec Épargne, définir un forfait facultatif par participante. Sans forfait, saisir le montant choisi lors de chaque cotisation.')
bullet('Enregistrer les cotisations d’une membre et consulter les tableaux de bord de campagne.')
bullet('Activer ou désactiver une campagne ; permettre aux membres de consulter les campagnes de leur groupement et leurs montants versés.')
h('Répartition des responsabilités')
table(['Rôle et interface','Droits prévus'],[
('Super administrateur\nApplication web','Créer les groupements et types de campagne ; ajouter des membres ; consulter les campagnes en lecture seule ; les activer ou les désactiver.'),
('Administrateur du groupement\nPWA','Créer et paramétrer les campagnes ; ajouter des membres ; enregistrer leurs cotisations ; consulter les tableaux de bord. Il conserve ses accès de membre.'),
('Membre\nPWA','Consulter les campagnes de son groupement et les montants versés. La restriction aux seuls versements personnels est proposée et reste à confirmer.')],[5,12])
p('Le super administrateur n’obtient pas automatiquement de droit de modification des cotisations. Son droit d’activation ou de désactivation constitue une permission distincte de la consultation.')
h('Hors du périmètre confirmé')
p('Paiement électronique, prêts, redistribution de l’épargne, notifications, exports métier, synchronisation des cotisations hors connexion et règles des autres types de campagne. Ces sujets pourront faire l’objet d’incréments ultérieurs. La traçabilité, les sauvegardes et la protection des données sont des fondations techniques proposées pour le premier incrément.')

page('2 Modèle métier et règles à valider')
h('Objets principaux proposés')
table(['Objet','Responsabilité'],[
('Compte utilisateur','Identité de connexion ; distincte de l’adhésion pour permettre l’ajout de membres sans compte actif.'),
('Groupement et adhésion','Informations du groupement, membre rattachée, statut et responsabilités locales.'),
('Affectation de rôle','Permissions d’administration limitées aux groupements attribués.'),
('Type de campagne','Catalogue global et version des règles prises en charge.'),
('Campagne','Groupement, type, période, calendrier, paramètres et état.'),
('Participation','Membre inscrite dans une campagne et forfait facultatif, avec historique des changements.'),
('Échéance','Date attendue selon le calendrier ; utile au suivi des cotisations, sous réserve des règles de paiement.'),
('Cotisation et audit','Montant réel, date du versement, date de saisie, auteur, référence et historique des corrections.')],[4,13])
h('Principes de fiabilité proposés')
p('Les montants utilisent une représentation exacte et une devise explicite ; le FCFA XOF est proposé, à confirmer. La date de versement est séparée de la date de saisie. Une cotisation est enregistrée dans une transaction et une nouvelle tentative du même enregistrement ne crée pas de doublon.')
p('Une correction conserve la trace de l’original, du motif et de son auteur. Les forfaits et paramètres historiques ne sont pas réécrits par un changement ultérieur. Les totaux sont calculés à partir des opérations valides et rapprochés avec les données sources.')
h('Décisions métier prioritaires avant développement')
bullet('Une membre peut-elle appartenir à plusieurs groupements ? Un administrateur peut-il gérer plusieurs groupements ? Qui lui attribue et retire ce rôle ?')
bullet('Comment activer les comptes, identifier les membres et récupérer un accès ? Le téléphone est-il obligatoire ou partagé ?')
bullet('Avec un forfait, autorise-t-on les versements partiels, les dépassements, les retards et plusieurs versements pour une échéance ?')
bullet('Qui peut corriger une cotisation ? Une campagne désactivée reste-t-elle consultable et peut-elle recevoir des corrections ?')
bullet('Quels indicateurs sont attendus ? Proposition : total versé, participantes, historique, versements par membre ; attendu et reste à cotiser uniquement lorsque le forfait permet de les calculer.')

page('3 Architecture et versionnement proposés')
h('Stack recommandée')
p('TypeScript en mode strict pour les deux interfaces et le backend ; React avec Vite pour le web et la PWA ; NestJS sur Node.js pour une API REST ; PostgreSQL pour les données relationnelles. La PWA utilise un manifest et un service worker. L’enregistrement des cotisations exige une connexion dans la proposition initiale.')
h('Décisions architecture à formaliser')
table(['ADR','Décision proposée'],[
('001','Monolithe modulaire : groupements, membres, campagnes, cotisations, accès et tableaux de bord.'),
('002','Base partagée avec isolation par groupement, contrôles backend et politiques PostgreSQL par ligne ; compte applicatif sans contournement.'),
('003','PostgreSQL avec contraintes, index et pagination ; partitionnement après mesure.'),
('004','Séparer compte, adhésion, participation et versement ; dater les changements de forfait.'),
('005','Cotisations transactionnelles, références uniques, idempotence et corrections traçables.'),
('006','Permissions par action et par groupement, vérifiées côté serveur.'),
('007','Versionner les règles des types et conserver celles appliquées aux campagnes.'),
('008','Tableaux de bord avec requêtes indexées puis agrégats si nécessaire ; traitements lourds en arrière-plan.'),
('009','API extensible sur plusieurs instances, sauvegardes, supervision et restauration testée.'),
('010','Conteneuriser les applications avec Docker ; Compose en développement et recette ; images immuables promues en production.')],[1.5,15.5])
h('Code et livraisons')
p('Un monorepo contient apps/super-admin-web, apps/membres-pwa et apps/api, ainsi que les composants UI, le client API généré depuis OpenAPI, les migrations et les ADR. La branche main est protégée ; les branches courtes sont fusionnées par pull request après les contrôles automatisés.')
p('Une version produit commune démarre à 0.1.0 pour la livraison interne, puis 1.0.0-beta.1 pour le pilote et 1.0.0 pour le premier incrément validé. Les corrections compatibles augmentent le correctif, les ajouts compatibles la version mineure et les ruptures de contrat la majeure. Un incrément métier supplémentaire peut donc être livré en 1.1.0.')
p('L’API commence à /api/v1. Les clients PWA supportés restent compatibles pendant une période définie. Les migrations sont numérotées et immuables après application ; les changements suivent ajout, migration, transition puis retrait. Chaque livraison possède un tag Git, un changelog et un artefact identique promu de recette vers production.')

page('4 Conteneurisation et déploiement')
p('La conteneurisation est intégrée au premier incrément. Docker est proposé pour construire et exécuter les services avec leurs dépendances. La PWA reste exécutée dans le navigateur ; son conteneur sert les fichiers compilés. Docker Compose décrit les services du développement local et de la recette.')
h('Services et données')
table(['Service','Conditionnement proposé'],[
('Web super administrateur','Image distincte servant les fichiers React compilés avec un serveur HTTP.'),
('PWA membres','Image distincte servant les fichiers compilés, le manifest et le service worker.'),
('API NestJS','Image Node.js avec les seules dépendances nécessaires à l’exécution.'),
('Traitements en arrière-plan','Si nécessaires, même image métier que l’API, avec commande et ressources distinctes.'),
('PostgreSQL','Conteneur et volume persistant en local et recette ; service PostgreSQL managé recommandé en production.'),
('Entrée HTTPS','Reverse proxy ou service d’entrée de la plateforme ; API et base restent sur des réseaux internes.')],[4.5,12.5])
h('Construction et configuration')
p('Utiliser des Dockerfiles à plusieurs étapes : installation, compilation puis exécution. Fixer les versions et empreintes des images de base, utiliser les fichiers de verrouillage des dépendances et un .dockerignore. Exécuter les applications sans privilèges administrateur ; injecter les secrets au déploiement, sans les copier dans les images ni dans Git.')
p('Pour promouvoir les mêmes images frontend entre environnements, charger la configuration publique, telle que l’URL de l’API, au démarrage ou utiliser une URL relative. Aucun secret ne doit figurer dans cette configuration. Prévoir contrôles de santé, limites CPU et mémoire, arrêt propre et logs vers la plateforme de supervision.')
h('Chaîne de livraison et versionnement')
p('La CI teste, construit et analyse les images, puis les publie dans un registre privé. Chaque image reçoit un tag de version produit et une référence du commit. Un manifeste de livraison fixe les empreintes exactes des images ; recette et production utilisent ces mêmes empreintes, sans reconstruction et sans tag latest.')
p('Les migrations sont exécutées une seule fois par un job contrôlé, avant le basculement selon leur compatibilité. Le retour arrière réutilise les anciennes images ; il ne supprime pas les données et suppose des migrations compatibles. Tester un redémarrage des conteneurs et vérifier la persistance des cotisations. Un volume ne remplace pas une sauvegarde.')
h('Choix de la plateforme de production')
p('Recommandation : plateforme de conteneurs managée avec plusieurs instances et PostgreSQL managé. Compose sur un serveur peut convenir au pilote si son indisponibilité est acceptée ; il ne fournit pas de haute disponibilité entre serveurs. Kubernetes reste une option ultérieure. L’ADR 010 précisera plateforme, registre, architectures CPU, stockage, secrets et budget. Conserver Dockerfiles, Compose et procédures dans le dépôt.')

page('5 Plan de réalisation du premier incrément')
p('Séquence indicative de 12 semaines pour deux développeurs à temps plein, avec un appui UX, QA et exploitation et un référent métier disponible. Cette estimation doit être révisée après cadrage ; elle suppose un fonctionnement connecté et aucune intégration de paiement.')
h('Étape 1 Cadrage et parcours utilisateurs  Semaine 1')
p('Résoudre les questions métier prioritaires, définir les permissions et indicateurs, préparer les maquettes des parcours web et PWA. Formaliser les ADR et le backlog ordonné. Responsable principal : référent métier avec équipe technique et UX.')
p('Validation : périmètre, règles de cotisation, stratégie des comptes et parcours approuvés ; critères d’acceptation associés aux fonctionnalités.')
h('Étape 2 Socle technique et sécurité  Semaines 2 et 3')
p('Créer monorepo, Dockerfiles, Compose, registre et CI de construction des images ; configurer authentification, rôles, isolation, schéma et migrations. Publier OpenAPI et mettre en place logs et sauvegardes. Responsable : équipe technique avec exploitation.')
p('Validation : démarrage Compose documenté ; connexion sur les deux interfaces ; accès interdits bloqués ; migration d’une base vierge et tests d’isolation réussis.')
h('Étape 3 Groupements et membres  Semaines 4 et 5')
p('Livrer la création des groupements dans le web, l’ajout et la consultation des membres, l’attribution contrôlée des administrateurs et la PWA installable. Distinguer création d’une adhésion et activation d’un compte.')
p('Validation : parcours complet de création d’un groupement et accès d’une administratrice à ses seuls groupements ; listes paginées et utilisables sur téléphone.')
h('Étape 4 Campagnes et cotisations  Semaines 6 à 8')
p('Livrer le catalogue des types, les campagnes Avec, le calendrier, les participations, les forfaits facultatifs et la saisie des cotisations. Ajouter idempotence, traçabilité, correction autorisée et activation ou désactivation selon les règles validées.')
p('Validation : scénarios avec et sans forfait, limites de période, double envoi et accès concurrents ; une coupure réseau ne donne jamais une confirmation de versement non enregistré.')
h('Étape 5 Tableaux de bord et qualification  Semaines 9 et 10')
p('Livrer les vues membre et administrateur, les consultations globales, les rapprochements de totaux, les tests de charge et les tests de restauration. Vérifier mise à jour de la PWA, confidentialité des caches et déconnexion sur téléphone partagé.')
p('Validation : totaux exacts, permissions conformes, charge et restauration validées ; redémarrage sans perte et analyse des images vérifiés ; aucune anomalie critique ouverte.')
h('Étape 6 Pilote et mise en production  Semaines 11 et 12')
p('Former les administratrices de 5 à 10 groupements pilotes, observer un cycle réel de collecte, corriger les anomalies et préparer les procédures de support. Livrer 1.0.0 après recette puis élargir par vagues, par exemple 50, 200 puis plus de 1 000 groupements.')
p('Validation : rapprochement des cotisations du pilote, accord métier, supervision active et procédure de retour à la version précédente testée. Les vagues suivantes dépendent des résultats observés, au-delà de cette estimation.')

page('6 Validation et conditions de lancement')
h('Scénarios de recette indispensables')
bullet('Le super administrateur crée un groupement et consulte ses campagnes sans pouvoir modifier les cotisations ; il peut changer leur activation.')
bullet('Une administratrice ajoute une membre, crée une campagne et enregistre un versement forfaitaire ou libre ; elle reste membre avec sa vue personnelle.')
bullet('Une utilisatrice d’un autre groupement ne peut accéder aux données, même avec une URL ou un identifiant modifié.')
bullet('Deux tentatives de la même saisie produisent une seule cotisation ; une correction reste retraçable et les totaux correspondent aux opérations valides.')
bullet('Une ancienne PWA supportée fonctionne avec l’API ; une mise à jour conserve les saisies en cours ou demande leur sauvegarde avant rechargement.')
h('Objectifs techniques à valider au cadrage')
p('Proposition initiale : temps API au 95e percentile inférieur à une seconde pour les lectures courantes et à deux secondes pour la saisie d’une cotisation, hors réseau utilisateur. Le chargement complet des écrans est mesuré séparément sur une connexion mobile représentative.')
p('Tester au moins 200 000 adhésions et environ 10 millions de cotisations, puis un scénario au-dessus de la cible. Définir la concurrence d’après les jours de collecte ; utiliser provisoirement 500 sessions actives avec un profil explicite de requêtes. Ces nombres sont des hypothèses de test et ne prouvent pas la capacité sans exécution et mesures.')
p('Proposer une perte de données maximale de 15 minutes et une restauration en moins de quatre heures, sous réserve de budget et d’infrastructure. Choisir la durée de conservation, les personnes habilitées et les procédures de sauvegarde et de restauration avant le pilote.')
h('Livrables de fin d’incrément')
bullet('Web, PWA, API et migrations ; Dockerfiles, configurations Compose, images versionnées, manifeste de livraison et procédure de déploiement.')
bullet('ADR validés, matrice des permissions, guide des administratrices et procédure de support.')
bullet('Rapports de recette, charge et restauration ; tableau de supervision ; journal des versions et procédure de déploiement.')
h('Conditions de passage en production')
p('Les parcours métier sont validés, les accès entre groupements sont isolés, les cotisations et totaux sont fiables, les sauvegardes sont restaurables et aucune anomalie critique n’est ouverte. Le pilote doit confirmer ces résultats avant extension du déploiement.')
page('7 Références techniques')
p('Ces documentations officielles précisent les mécanismes sur lesquels reposent les propositions techniques. Les choix de plateforme et les versions des dépendances seront figés au démarrage du développement.')
p('TypeScript : https://www.typescriptlang.org/docs/handbook/intro\nNestJS : https://docs.nestjs.com/first-steps\nReact : https://react.dev/learn/typescript\nVite : https://vite.dev/guide/\nPostgreSQL sécurité : https://www.postgresql.org/docs/current/ddl-rowsecurity.html\nPostgreSQL partitionnement : https://www.postgresql.org/docs/current/ddl-partitioning.html\nVersionnement sémantique : https://semver.org/')
p('Docker Compose : https://docs.docker.com/compose/\nDocker Compose en production : https://docs.docker.com/compose/how-tos/production/\nDocker bonnes pratiques : https://docs.docker.com/build/building/best-practices/\nDocker construction en plusieurs étapes : https://docs.docker.com/build/building/multi-stage/')
doc.save(ROOT / 'AvecLesFemmes_Resume_et_Plan_Implementation.docx')
print('Document créé')
